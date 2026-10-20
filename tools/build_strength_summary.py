"""
ルールAI(AIルール相談機能)向けの「強さの目安」要約データを生成する。

C:\\Users\\miwa\\Desktop\\質問AI用エクセル\\ フォルダに置かれた、ai_data_report.js/ai_data_write.py
形式のAIDATAレポート(*.xlsx)から、JOB/CONそれぞれの「平均得点・QST平均得点・平均順位」という
軽量な要約(CONJOBシートの31/51/69行目=JOB側、O列18-29/38-49/56-67行目=CON側)だけを取り出し、
data/ai-strength-summary.json に書き出す。

フォルダ内に複数の.xlsxがある場合は、更新日時が一番新しいものを自動で使う(2026-10-21、per user
request: "最新版に更新できるように" -- ユーザーが新しいレポートをそのフォルダに置き換えるだけで、
このスクリプトを再実行すれば常に最新のAI対戦結果を反映できる)。

評価値シート(AIの内部判断用の重み)ではなく、実際にAI同士が対戦した結果(勝敗・順位)から来る
この要約を使うのは、人間のプレイヤーにとって「このJOBは平均順位が良い」の方が、内部的な重み
の数値よりずっと分かりやすいため(2026-10-21、per user提案)。

Usage: python tools/build_strength_summary.py
"""
import glob
import json
import os
import sys
from datetime import datetime, timezone

import openpyxl

SOURCE_DIR = r'C:\Users\miwa\Desktop\質問AI用エクセル'
OUTPUT_PATH = os.path.join(os.path.dirname(__file__), '..', 'data', 'ai-strength-summary.json')

JOB_SCORE_ROW = 31
JOB_QST_SCORE_ROW = 51
JOB_RANK_ROW = 69
JOB_NAME_ROW = 1
JOB_FIRST_COL = 2  # column B

CON_NAME_ROWS = list(range(2, 14))  # rows 2-13
CON_SCORE_COL = 15  # column O
CON_SCORE_ROW_OFFSET = 16  # CON name row N -> score at row N+16 (2->18, 13->29)
CON_QST_SCORE_ROW_OFFSET = 36  # 2->38, 13->49
CON_RANK_ROW_OFFSET = 54  # 2->56, 13->67


def find_latest_source_file():
    candidates = glob.glob(os.path.join(SOURCE_DIR, '*.xlsx'))
    candidates = [c for c in candidates if not os.path.basename(c).startswith('~$')]  # Excelのロックファイル除外
    if not candidates:
        raise FileNotFoundError(f'{SOURCE_DIR} に.xlsxファイルが見つかりません')
    return max(candidates, key=os.path.getmtime)


def round_or_none(value):
    return round(value, 2) if isinstance(value, (int, float)) else None


def build_summary(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['CONJOB']

    job = {}
    col = JOB_FIRST_COL
    while True:
        name = ws.cell(row=JOB_NAME_ROW, column=col).value
        if not name:
            break
        job[name] = {
            'avgScore': round_or_none(ws.cell(row=JOB_SCORE_ROW, column=col).value),
            'avgQstScore': round_or_none(ws.cell(row=JOB_QST_SCORE_ROW, column=col).value),
            'avgRank': round_or_none(ws.cell(row=JOB_RANK_ROW, column=col).value),
        }
        col += 1

    con = {}
    for name_row in CON_NAME_ROWS:
        name = ws.cell(row=name_row, column=1).value
        if not name:
            continue
        con[name] = {
            'avgScore': round_or_none(ws.cell(row=name_row + CON_SCORE_ROW_OFFSET, column=CON_SCORE_COL).value),
            'avgQstScore': round_or_none(ws.cell(row=name_row + CON_QST_SCORE_ROW_OFFSET, column=CON_SCORE_COL).value),
            'avgRank': round_or_none(ws.cell(row=name_row + CON_RANK_ROW_OFFSET, column=CON_SCORE_COL).value),
        }

    return {
        '_comment': 'avgScore/avgQstScore=平均得点(高いほど良い)、avgRank=平均順位(低いほど良い、1位に近いほど良い)。JOB/CONはそれぞれ単体の平均で、組み合わせ(例: 特定JOB×特定CON)ごとの数値ではない。',
        'sourceFile': os.path.basename(path),
        'generatedAt': datetime.now(timezone.utc).isoformat(),
        'job': job,
        'con': con,
    }


def main():
    source_path = find_latest_source_file()
    summary = build_summary(source_path)
    with open(OUTPUT_PATH, 'w', encoding='utf-8') as f:
        json.dump(summary, f, ensure_ascii=False, indent=1)
    print(f'Wrote {OUTPUT_PATH} (source: {source_path}, {len(summary["job"])} JOBs, {len(summary["con"])} CONs)')


if __name__ == '__main__':
    main()
