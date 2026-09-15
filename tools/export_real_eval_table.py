"""
Exports the REAL (currently live, hand-tuned) 評価値 table -- i.e. exactly what AI LV4 actually reads
at runtime, since LV4 has no separate trained genome wired into production (see src/ai/levels.js) --
into a standalone .xlsx, using the same NAME-lookup + game.xlsx's own row order as the eval-table block
in tools/ga_progress_to_xlsx.py (2026-09-15, per user request to export "現在のAI LV4の評価値" in the
same style, after an earlier plain-ID/alphabetical attempt via ga_genome_to_xlsx.py was rejected for
being "ばらばら" and ID-labeled instead of NAME-labeled).

export(xlsx_path) is also called automatically from tools/xlsx_to_json.py after every game.xlsx ->
game.json regeneration (2026-09-15, per user request: "AILV4の評価値が変わるごとにここに書いてある方式
で記入お願い" -- game.xlsx's 評価値 sheet IS AI LV4's real eval table, so any edit to it that goes
through the normal xlsx_to_json.py rebuild step should also refresh this standalone snapshot, at a
fixed path the user keeps open/reuses, without a separate manual step).

Usage (standalone):
    python tools/export_real_eval_table.py <outputXlsxPath>
"""

import json
import os
import sys
import openpyxl

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
GAME_JSON_PATH = os.path.join(SCRIPT_DIR, '..', 'data', 'game.json')

# Fixed destination the user asked to keep in sync automatically -- see this module's own doc.
DEFAULT_XLSX_PATH = r'C:\Users\miwa\Desktop\AILV4_評価値_20260914.xlsx'


def export(xlsx_path, game_json_path=GAME_JSON_PATH):
    with open(game_json_path, 'r', encoding='utf-8') as f:
        game_data = json.load(f)

    # ID -> NAME lookup (2026-09-12 convention, see ga_progress_to_xlsx.py's load_id_to_name): covers
    # every card-face row across A/B/C/M; plain resource ids and already-Japanese synergy-bonus row
    # names have no match and fall back to their own id/label unchanged.
    id_to_name = {}
    for sheet in ('A', 'B', 'C', 'M'):
        for row in game_data.get(sheet, []):
            if row.get('ID') and row.get('NAME'):
                id_to_name[row['ID']] = row['NAME']

    rows = game_data.get('評価値', [])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '評価値'

    ws.cell(row=1, column=1, value='ID')
    for i, r in enumerate((1, 2, 3, 4)):
        ws.cell(row=1, column=2 + i, value=f'{r}R')

    row_idx = 2
    for row in rows:
        id_ = row['ID']
        label = id_to_name.get(id_, id_)
        ws.cell(row=row_idx, column=1, value=label)
        for i, r in enumerate((1, 2, 3, 4)):
            raw = row.get(f'{r}R')
            value = raw if isinstance(raw, (int, float)) else 0
            ws.cell(row=row_idx, column=2 + i, value=round(value, 2) if isinstance(value, float) else value)
        row_idx += 1

    ws.column_dimensions['A'].width = 24
    for col_letter in ('B', 'C', 'D', 'E'):
        ws.column_dimensions[col_letter].width = 14

    wb.save(xlsx_path)
    return len(rows)


if __name__ == '__main__':
    xlsx_path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX_PATH
    count = export(xlsx_path)
    print(f'Wrote {xlsx_path} ({count} rows, game.xlsx order, NAME-labeled)')
