# -*- coding: utf-8 -*-
"""
One-shot script (2026-09-07, per user request to add card-ownership-synergy-aware eval rows --
see src/ai/evaluator.js's new SYNERGY_ROW_NAMES block for how these are actually scored):
adds 2 new GA-tunable 評価値 rows, seeded at 0 in every round like モニュメント確保ボーナス before them
(src/ai/ga.js's mutateGenomePercent "zero-escape" step lets GA/replay training discover a nonzero
value for these from scratch, same mechanism that row already uses).

  - 聖女王女アンタップ相性: credited once per owned untap-worthy B/C card (see UNTAP_SYNERGY_FACE_IDS in
    evaluator.js) while the player owns 聖女 or 王女 (C202/C301, either tier) -- per user: "タップする
    カードで毎ターン使えないカード(兆し/導き/Cカード)があると聖女・王女が強い".
  - 晩餐会食料生産相性: credited once while the player owns 農園/小麦畑/農夫 (A004/A005/C201, either
    tier) and 晩餐会(M401) is still unclaimed -- per user: "農園 小麦畑 農民などがあると...晩餐会...が
    狙いやすい".

Run once: `python tools/add_synergy_eval_rows.py`, then `python tools/xlsx_to_json.py` to recompile.
"""
import openpyxl

PATH = 'data/game.xlsx'
NEW_ROWS = ['聖女王女アンタップ相性', '晩餐会食料生産相性']

wb = openpyxl.load_workbook(PATH)
ws = wb['評価値']

existing_ids = {ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)}
last_real = 0
for r in range(1, ws.max_row + 1):
    if ws.cell(row=r, column=1).value is not None:
        last_real = r

next_row = last_real + 1
for name in NEW_ROWS:
    if name in existing_ids:
        print(f'Skipping {name!r} -- already present')
        continue
    ws.cell(row=next_row, column=1, value=name)
    for col, _round in zip(range(2, 6), ['1R', '2R', '3R', '4R']):
        ws.cell(row=next_row, column=col, value=0)
    print(f'Added row {next_row}: {name}')
    next_row += 1

wb.save(PATH)
print('Saved', PATH)
