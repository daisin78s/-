"""
Writes tools/replay_round_end_scores.js's own output JSON into an .xlsx workbook -- one sheet per
replay ("リプレイ1".."リプレイ13"), each sheet stacking its own 1R/2R/3R/4R blocks vertically
(2026-09-12, per user request: "リプレイ1の1R 2R 3R 4Rという風にエクセル出力で").

Each round block is a small table: one column per player (name, HUMAN marker if applicable), one row
per scoring line item (K/A/B/C/VP/dice terms + "その他(カード/シナジー等)"), plus a 合計(total) row.

Usage: python tools/replay_round_end_scores_write_xlsx.py <inputJsonPath> <outputXlsxPath>
"""

import json
import sys
import openpyxl
from openpyxl.styles import Font

if len(sys.argv) < 3:
    print('Usage: python tools/replay_round_end_scores_write_xlsx.py <inputJsonPath> <outputXlsxPath>')
    sys.exit(1)

INPUT_PATH = sys.argv[1]
OUTPUT_PATH = sys.argv[2]

with open(INPUT_PATH, 'r', encoding='utf-8') as f:
    data = json.load(f)

wb = openpyxl.Workbook()
wb.remove(wb.active)

BOLD = Font(bold=True)

for replay in data:
    ws = wb.create_sheet(replay['label'])
    row = 1
    for round_block in replay['rounds']:
        ws.cell(row=row, column=1, value=f"{round_block['round']}R終了時").font = BOLD
        row += 1

        players = round_block['players']
        # Header row: player names (blank first column reserved for item labels).
        for col, player in enumerate(players, start=2):
            label = player['name'] or player['playerId']
            if player['isHuman']:
                label += '(人間)'
            ws.cell(row=row, column=col, value=label).font = BOLD
        row += 1

        # Union of every item label across this round's players, in first-seen order, so a player with
        # a zero/absent term (filtered out upstream) still lines up under the right row.
        item_labels = []
        for player in players:
            for item in player['items']:
                if item['label'] not in item_labels:
                    item_labels.append(item['label'])

        for label in item_labels:
            ws.cell(row=row, column=1, value=label)
            for col, player in enumerate(players, start=2):
                match = next((it for it in player['items'] if it['label'] == label), None)
                if match is None:
                    continue
                if match['count'] is not None:
                    ws.cell(row=row, column=col, value=f"{match['count']}*{match['weight']}={round(match['contribution'], 1)}")
                else:
                    ws.cell(row=row, column=col, value=round(match['contribution'], 1))
            row += 1

        ws.cell(row=row, column=1, value='合計').font = BOLD
        for col, player in enumerate(players, start=2):
            ws.cell(row=row, column=col, value=round(player['total'], 1)).font = BOLD
        row += 2  # blank spacer row before the next round block

    ws.column_dimensions['A'].width = 26
    for col in range(2, 6):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 22

wb.save(OUTPUT_PATH)
print(f'Wrote {OUTPUT_PATH}')
