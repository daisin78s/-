"""
Writes カード相性.xlsx from tools/card_pair_stats.js's accumulated tally (see that file's own doc for the
request and for what one row/column "カード" is). Called by card_pair_stats.js's save(), same Node->Python
subprocess pattern ga_train.js already uses for ga_progress_to_xlsx.py.

Sheets:
  - 相性:     (予想の平均順位) - (実際の平均順位). Plus = the two cards together do better than each card's own
              solo average predicts (good together), minus = worse. Predicted = solo(a) + solo(b) - overall
              average, so a card that's simply strong on its own doesn't make every pair with it look "good" --
              only the extra on top of that shows. Blank until the pair has MIN_GAMES games.
  - 補正後の相性: 相性 minus the average 相性 of its pair type (area x area, monument x monument, job x
              area, ...; 2026-10-30, per user request "その計算でお願い"). Every pair type averages clearly above 0
              in raw 相性 (+0.08 area x area .. +0.15 monument x monument), because a player who ends up holding
              more cards is usually the one winning -- subtracting each type's own average removes that bias so
              pairs of different types can be compared on one scale. The type average uses pairs with at least
              BASELINE_MIN_GAMES games (fewer-game pairs are too noisy to set a baseline).
  - 補正後ランキング: every pair with at least BASELINE_MIN_GAMES games, sorted by 補正後の相性.
  - 平均順位: average final rank of a player who held both cards at game end (1 best .. 4 worst). The
              diagonal is that card's own solo average rank.
  - 試合数:   how many player-games each cell is based on.
  - 一覧:     each card's own solo numbers, plus the log of rows/columns deleted because an ability changed.

Usage: python tools/card_pair_stats_to_xlsx.py <statsJson> <gameJson> <xlsxOut>
"""

import json
import re
import sys
import openpyxl
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill

MIN_GAMES = 30
BASELINE_MIN_GAMES = 100
UNIT_SHEETS = ['A', 'B', 'C', 'M', 'JOB', 'CON']
KIND_LABEL = {'A': 'A', 'B': 'B', 'C': 'C', 'M': 'モニュメント', 'JOB': 'ジョブ', 'CON': '性格'}
# Pair-type grouping for 補正後の相性: A/B/C are all area cards, one group.
GROUP_LABEL = {'A': 'エリア', 'B': 'エリア', 'C': 'エリア', 'M': 'モニュメント', 'JOB': 'ジョブ', 'CON': '性格'}

stats_path, game_path, xlsx_path = sys.argv[1], sys.argv[2], sys.argv[3]
with open(stats_path, encoding='utf-8') as f:
    stats = json.load(f)
with open(game_path, encoding='utf-8') as f:
    game = json.load(f)

# Same unit rule as card_pair_stats.js's unitIdOfFace (A/B/C physical card; M/JOB/CON face id as-is).
units = []
seen = set()
for sheet in UNIT_SHEETS:
    for row in game.get(sheet, []):
        uid = row['ID'] if sheet in ('M', 'JOB', 'CON') else re.sub(r'[AB]$', '', row['ID'])
        if uid in seen:
            continue
        seen.add(uid)
        units.append((uid, sheet, re.sub(r'LV\d$', '', row.get('NAME') or uid)))

solo = stats.get('solo', {})
pair = stats.get('pair', {})
total_n = sum(s['n'] for s in solo.values())
overall = (sum(s['rankSum'] for s in solo.values()) / total_n) if total_n else 2.5


def solo_avg(u):
    s = solo.get(u)
    return s['rankSum'] / s['n'] if s and s['n'] else None


def pair_cell(a, b):
    return pair.get(f'{a}|{b}' if a < b else f'{b}|{a}')


wb = openpyxl.Workbook()
del wb['Sheet']
header_fill = PatternFill('solid', fgColor='EEEEEE')
diag_fill = PatternFill('solid', fgColor='D9D9D9')


def matrix_sheet(title, value_of, number_format, color_rule=None):
    ws = wb.create_sheet(title)
    ws.cell(row=1, column=1, value=f'{stats.get("games", 0)}試合分').font = Font(bold=True)
    for i, (uid, sheet, name) in enumerate(units):
        h = ws.cell(row=1, column=2 + i, value=name)
        h.fill = header_fill
        h.alignment = Alignment(text_rotation=90, horizontal='center')
        r = ws.cell(row=2 + i, column=1, value=name)
        r.fill = header_fill
    for i, (a, _, _) in enumerate(units):
        for j, (b, _, _) in enumerate(units):
            v = value_of(a, b, i == j)
            c = ws.cell(row=2 + i, column=2 + j, value=v)
            c.number_format = number_format
            if i == j:
                c.fill = diag_fill
    ws.freeze_panes = 'B2'
    ws.column_dimensions['A'].width = 18
    for j in range(len(units)):
        ws.column_dimensions[openpyxl.utils.get_column_letter(2 + j)].width = 6
    ws.row_dimensions[1].height = 110
    if color_rule:
        last = openpyxl.utils.get_column_letter(1 + len(units))
        ws.conditional_formatting.add(f'B2:{last}{1 + len(units)}', color_rule)
    return ws


def synergy(a, b, diag):
    if diag:
        return None
    p = pair_cell(a, b)
    sa, sb = solo_avg(a), solo_avg(b)
    if not p or p['n'] < MIN_GAMES or sa is None or sb is None:
        return None
    return round((sa + sb - overall) - p['rankSum'] / p['n'], 3)


def avg_rank(a, b, diag):
    if diag:
        v = solo_avg(a)
        return round(v, 3) if v is not None else None
    p = pair_cell(a, b)
    return round(p['rankSum'] / p['n'], 3) if p else None


def count(a, b, diag):
    if diag:
        s = solo.get(a)
        return s['n'] if s else None
    p = pair_cell(a, b)
    return p['n'] if p else None


group_of = {uid: GROUP_LABEL[sheet] for uid, sheet, _ in units}
name_of = {uid: name for uid, _, name in units}


def pair_type(a, b):
    return '×'.join(sorted([group_of[a], group_of[b]]))


baseline_sum = {}
baseline_n = {}
for i, (a, _, _) in enumerate(units):
    for b, _, _ in units[i + 1:]:
        p = pair_cell(a, b)
        v = synergy(a, b, False)
        if v is None or p['n'] < BASELINE_MIN_GAMES:
            continue
        t = pair_type(a, b)
        baseline_sum[t] = baseline_sum.get(t, 0) + v
        baseline_n[t] = baseline_n.get(t, 0) + 1
baseline = {t: baseline_sum[t] / baseline_n[t] for t in baseline_sum}


def adjusted_synergy(a, b, diag):
    v = synergy(a, b, diag)
    if v is None:
        return None
    return round(v - baseline.get(pair_type(a, b), 0), 3)


synergy_colors = dict(start_type='num', start_value=-0.5, start_color='F8696B',
                      mid_type='num', mid_value=0, mid_color='FFFFFF',
                      end_type='num', end_value=0.5, end_color='63BE7B')
matrix_sheet('補正後の相性', adjusted_synergy, '+0.00;-0.00;0.00', ColorScaleRule(**synergy_colors))

ws = wb.create_sheet('補正後ランキング')
ws.append(['順位', 'カード1', 'カード2', '組の種類', '試合数', '平均順位', '予想', '相性', '補正後の相性'])
ranking = []
for i, (a, _, _) in enumerate(units):
    for b, _, _ in units[i + 1:]:
        p = pair_cell(a, b)
        v = adjusted_synergy(a, b, False)
        if v is None or p['n'] < BASELINE_MIN_GAMES:
            continue
        ranking.append((v, a, b, p))
ranking.sort(key=lambda r: -r[0])
for k, (v, a, b, p) in enumerate(ranking, 1):
    ws.append([k, name_of[a], name_of[b], pair_type(a, b), p['n'], round(p['rankSum'] / p['n'], 3),
               round(solo_avg(a) + solo_avg(b) - overall, 3), synergy(a, b, False), v])
ws.freeze_panes = 'A2'
for col, width in zip('BCD', (18, 18, 16)):
    ws.column_dimensions[col].width = width

matrix_sheet('相性', synergy, '+0.00;-0.00;0.00', ColorScaleRule(**synergy_colors))
matrix_sheet('平均順位', avg_rank, '0.00',
             ColorScaleRule(start_type='num', start_value=1.5, start_color='63BE7B',
                            mid_type='num', mid_value=2.5, mid_color='FFFFFF',
                            end_type='num', end_value=3.5, end_color='F8696B'))
matrix_sheet('試合数', count, '0')

ws = wb.create_sheet('一覧')
ws.append(['ID', '名前', '種類', '持っていた数', '平均順位'])
for uid, sheet, name in units:
    s = solo.get(uid)
    ws.append([uid, name, KIND_LABEL[sheet], s['n'] if s else 0, round(s['rankSum'] / s['n'], 3) if s and s['n'] else None])
ws.append([])
ws.append(['見方'])
ws.append([f'補正後の相性: 相性から、その組の種類(エリア×エリア、モニュメント×モニュメントなど)の平均を引いたもの。たくさんカードを持っている人ほど勝っているくせを消して、種類の違う組どうしを同じ基準で比べられるようにしたもの(平均は{BASELINE_MIN_GAMES}試合以上の組で計算)'])
ws.append([f'相性: 2枚を一緒に持っていた時の平均順位が、それぞれ単独の平均順位から予想される値より、どれだけ良いか(+ほど良い)。{MIN_GAMES}試合未満は空欄'])
ws.append(['平均順位: 2枚とも持っていたプレイヤーの最終順位の平均(1位が一番良い)。斜めのマスはそのカード単独の平均順位'])
ws.append(['数え方: 進化AIの対戦の終了時に持っていたカード(初期資源は除く)。A/B/CはLV1とLV2を同じカードとして数える。性格は選んだ面ごと'])
ws.append([])
ws.append(['組の種類ごとの平均(補正後の相性で引いている値)', '', '組数'])
for t in sorted(baseline, key=lambda t: -baseline[t]):
    ws.append([t, round(baseline[t], 3), baseline_n[t]])
ws.append([])
ws.append(['削除の履歴(能力が変わったカードは、その行と列の数字を消して0から数え直す)'])
for entry in stats.get('resetLog', []):
    ws.append([entry['unitId'], entry['reason'], entry['at']])
ws.column_dimensions['B'].width = 18

wb.save(xlsx_path)
