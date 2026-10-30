"""
Writes game.xlsx's 評価値ALL sheet (what the AI uses, see src/ai/eval-all.js) and the 確認用 tables to the right
of 評価値_JOB / 評価値_CON from the カード相性 data (tools/card_pair_stats.js). 2026-10-30, per user requests:
  - "評価値ALLシート作りました これにすべての組み合わせをいれて それを使う" + "c" (評価値_JOB/評価値_CON no longer
    used by the AI) + "500ごとに更新" (tools/eval_all_update.js calls this, tools/ga_hillclimb.js calls that).
  - "プラス効果は入力 マイナスは一部以外 入力しない マイナスを入力すると AIがとらなくなり 検証できなくなる".
  - "データをとるまでもなくマイナスは明らかなので-100くらいを入れておいて" -> fixed cells (below).
  - "評価値 JOBのシート 右側に もう一つ 表を作りました そこにALLでいれる評価値を 入れてください これも自動更新
    するように これには マイナスも すべて入れる これは確認用で AIには使わない 同じことを ほかの 評価値シート
    にもする" -> the 確認用 tables (every pair with data, minus included, no thresholds).

Points from 補正後の相性 (see card_pair_synergy.py): POINTS_PER_RANK x 補正後 x n/(n+SHRINK_GAMES) -- the
n/(n+SHRINK_GAMES) factor shrinks pairs with few games toward 0 (100 games -> 1/3, 1000 games -> 5/6), since
their 相性 is mostly noise. POINTS_PER_RANK = 130 lines a 0.15 / 0.30 rank edge up with the old hand-typed
+20 / +40 values in 評価値_JOB / 評価値_CON.

評価値ALL cells:
  - From data: 補正後 >= PLUS_MIN -> the points. 補正後 <= MINUS_MAX with at least MINUS_MIN_GAMES games -> the
    points, but never below -99 (so it can't be mistaken for a fixed cell). Anything else blank.
  - Fixed (never rewritten, mirrored to both triangles): any cell <= -100, and the whole D row/column (D is
    追加色ダイス, not a card, so there's no data for it). --seed-fixed copies them once from the old sheets:
    every old cell <= -100 becomes -100, and the old D rows are copied as-is.

Usage: python tools/write_eval_all_sheet.py <statsJson> <gameJson> <xlsx> [--seed-fixed]
"""

import json
import re
import sys
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, __import__('os').path.dirname(__file__))
from card_pair_synergy import PairSynergy  # noqa: E402

POINTS_PER_RANK = 130
SHRINK_GAMES = 200
PLUS_MIN = 0.10
MINUS_MAX = -0.30
MINUS_MIN_GAMES = 300
FIXED_MAX = -100
D = 'D'
SHEET = '評価値ALL'

args = [a for a in sys.argv[1:] if not a.startswith('--')]
seed_fixed = '--seed-fixed' in sys.argv
stats_path, game_path, xlsx_path = args
with open(stats_path, encoding='utf-8') as f:
    stats = json.load(f)
with open(game_path, encoding='utf-8') as f:
    game = json.load(f)

ps = PairSynergy(stats, game)
unit_by_name = {name: uid for uid, _, name in ps.units}
unit_by_name[D] = D
for row in game.get('CON', []):
    unit_by_name[row['NAME']] = row['ID']
for row in game.get('JOB', []):
    unit_by_name[row['ID']] = row['ID']  # 評価値_JOB's header row holds JOB ids


def unit_of(label):
    if label is None:
        return None
    return unit_by_name.get(re.sub(r'LV\d+$', '', str(label)))


def points(a, b):
    """(points, n, adjusted) for a pair, or None if there's no data yet."""
    if D in (a, b) or a == b:
        return None
    adj = ps.adjusted_synergy(a, b)
    if adj is None:
        return None
    n = ps.pair_cell(a, b)['n']
    return round(POINTS_PER_RANK * adj * n / (n + SHRINK_GAMES)), n, adj


def ai_points(a, b):
    p = points(a, b)
    if p is None:
        return None
    value, n, adj = p
    if adj >= PLUS_MIN and value > 0:
        return value
    if adj <= MINUS_MAX and n >= MINUS_MIN_GAMES and value < 0:
        return max(value, FIXED_MAX + 1)
    return None


wb = openpyxl.load_workbook(xlsx_path)

# --- fixed cells ---
fixed = {}  # (a, b) -> value, both orders


def set_fixed(a, b, v):
    fixed[(a, b)] = v
    fixed[(b, a)] = v


ws = wb[SHEET] if SHEET in wb.sheetnames else wb.create_sheet(SHEET)
if ws.max_row >= 2:
    header = [c.value for c in ws[1]]
    for r in ws.iter_rows(min_row=2):
        a = unit_of(r[0].value)
        if not a:
            continue
        for idx, cell in enumerate(r[1:], start=1):
            b = unit_of(header[idx]) if idx < len(header) else None
            if not b or b == a or not isinstance(cell.value, (int, float)):
                continue
            if cell.value <= FIXED_MAX or D in (a, b):
                set_fixed(a, b, cell.value)

if seed_fixed:
    for sheet_name, header_row in (('評価値_JOB', 1), ('評価値_CON', 1)):
        src = wb[sheet_name]
        header = [c.value for c in src[header_row]]
        for r in src.iter_rows(min_row=header_row + 1):
            a = unit_of(r[0].value)
            if not a:
                continue
            for idx in range(1, len(header)):
                if header[idx] is None:
                    break  # the left (real) table ends at the first blank header column
                b = unit_of(header[idx])
                v = r[idx].value if idx < len(r) else None
                if not b or b == a or not isinstance(v, (int, float)):
                    continue
                if v <= FIXED_MAX:
                    set_fixed(a, b, FIXED_MAX)
                elif D in (a, b):
                    set_fixed(a, b, v)

# --- 評価値ALL ---
order = [uid for uid, _, _ in ps.units] + [D]
label = {uid: name for uid, _, name in ps.units}
label[D] = D
ws.delete_rows(1, ws.max_row)
header_fill = PatternFill('solid', fgColor='EEEEEE')
fixed_font = Font(bold=True, color='C00000')
ws.cell(row=1, column=1, value='NAME').fill = header_fill
for j, b in enumerate(order):
    c = ws.cell(row=1, column=2 + j, value=label[b])
    c.fill = header_fill
    c.alignment = Alignment(text_rotation=90, horizontal='center')
written = 0
for i, a in enumerate(order):
    ws.cell(row=2 + i, column=1, value=label[a]).fill = header_fill
    for j, b in enumerate(order):
        if a == b:
            continue
        if (a, b) in fixed:
            c = ws.cell(row=2 + i, column=2 + j, value=fixed[(a, b)])
            c.font = fixed_font
            continue
        v = ai_points(a, b)
        if v is not None:
            ws.cell(row=2 + i, column=2 + j, value=v)
            written += 1
ws.freeze_panes = 'B2'
ws.column_dimensions['A'].width = 18
for j in range(len(order)):
    ws.column_dimensions[openpyxl.utils.get_column_letter(2 + j)].width = 5
ws.row_dimensions[1].height = 110


# --- 確認用 tables right of 評価値_JOB / 評価値_CON ---
def write_check_table(sheet_name):
    src = wb[sheet_name]
    header = [c.value for c in src[1]]
    left_width = next((i for i, v in enumerate(header) if v is None), len(header))
    start = left_width + 1  # 0-based column index of the 確認用 table (one blank column in between)
    # Header + row labels: the same as the left table (created here if the sheet doesn't have one yet).
    for idx in range(left_width):
        src.cell(row=1, column=start + idx + 1, value=header[idx])
    count = 0
    for r in range(2, src.max_row + 1):
        row_label = src.cell(row=r, column=1).value
        src.cell(row=r, column=start + 1, value=row_label)
        a = unit_of(row_label)
        for idx in range(1, left_width):
            b = unit_of(header[idx])
            p = points(a, b) if a and b else None
            v = p[0] if p and p[0] != 0 else None
            if (a, b) in fixed:
                v = fixed[(a, b)]
            if v is None and row_label in (None, 'NAME'):
                v = src.cell(row=r, column=idx + 1).value  # 評価値_JOB's JOB-name legend row
            # .value directly: ws.cell(..., value=None) leaves the old value in place, which kept the
            # hand values copied from the left table showing in cells that have no data yet.
            src.cell(row=r, column=start + idx + 1).value = v
            if p:
                count += 1
    return count


check_counts = {name: write_check_table(name) for name in ('評価値_JOB', '評価値_CON')}
wb.save(xlsx_path)
print(json.dumps({'games': stats.get('games', 0), 'evalAllCells': written, 'fixedCells': len(fixed) // 2, 'check': check_counts}, ensure_ascii=False))
