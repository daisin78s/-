"""Convert chapter 2's script (ストーリー.xlsx, sheet "Sheet2") into data/story2.data.js.

Usage: python tools/story2_xlsx_to_js.py [path/to/ストーリー.xlsx]
Input: the argument, otherwise ~/Desktop/ストーリー.xlsx. Sheet "Sheet2" (or the 2nd sheet if renamed).
Output: data/story2.data.js, a plain browser global (window.STORY2_SCRIPT) loaded via a normal <script src>
tag -- same reasoning as data/game.data.js (index.html must keep working from file://).

2026-10-31, per user: chapter 2 (○○○の百人隊長スパル) is written in Excel instead of Twine ("エクセルのほうです"):
one row per step, in order, no IDs -- column A is a line (spoken by スパル), column C is what happens on the
screen or the condition that starts that part ("ダイスが振られる", "訓練場にダイスを置く（１回目）", ...). Both
columns are exported as-is ({text, event}); story.js decides what each event means by its wording (see
story.js's chapter 2 section). Dev-time tool only; the game never reads the xlsx directly.
"""
import json
import sys
from pathlib import Path

import openpyxl

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DEST_JS = PROJECT_ROOT / "data" / "story2.data.js"
DEFAULT_XLSX = Path.home() / "Desktop" / "ストーリー.xlsx"
SHEET_NAME = "Sheet2"


def cell_text(value):
    return "" if value is None else str(value).strip()


def main():
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    wb = openpyxl.load_workbook(src, data_only=True)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.worksheets[1]
    rows = []
    for row in ws.iter_rows(min_row=1, max_col=3, values_only=True):
        text, event = cell_text(row[0]), cell_text(row[2] if len(row) > 2 else None)
        if text or event:
            rows.append({"text": text, "event": event})
    inline = json.dumps(rows, ensure_ascii=False, indent=1).replace("</script", "<\\/script")
    DEST_JS.write_text(f"window.STORY2_SCRIPT = {inline};\n", encoding="utf-8")
    print(f"{len(rows)} 行 <- {src} ({ws.title})")
    print(f"-> {DEST_JS}")


if __name__ == "__main__":
    main()
